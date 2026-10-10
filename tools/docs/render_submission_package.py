#!/usr/bin/env python3
"""Build the CIT-U IT411 submission package in docs/submission/.

- SDD v2.0, SRS v3.0 and SPMP v2.0 are rendered from their Markdown sources
  (md_to_docx.py), each with a cover page, a table of contents, page numbers,
  and landscape pages for the wide matrices.
- The MVP Validation Findings & Refactoring Priorities form is the course
  template filled from its Markdown source (fill_mvp_validation_form.py).
- Microsoft Word converts every .docx to .pdf and fills in the tables of
  contents (PowerShell COM, so it runs from Windows or from WSL).

Usage: python3 tools/docs/render_submission_package.py   (needs python-docx and Word)
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
except ImportError:
    sys.exit("Error: python-docx is required. Run: pip install python-docx")

import fill_mvp_validation_form
from md_to_docx import (BODY_FONT, HEADING_SIZES, MUTED, PLUM, TEAL, TEXT, Theme, add_field, add_inline,
                        page_break, parse_blocks, render_blocks, set_run_font)

REPO = Path(__file__).resolve().parents[2]
DOCS = REPO / "docs"
SUBMISSION_DIR = DOCS / "submission"

SPEC_DOCUMENTS = [
    {
        "src": DOCS / "SDD_v2.0_Refactored.md",
        "out": "playIT_SDD_v2.0",
        "title": "Software Design Description (SDD v2.0)",
        "short": "SDD v2.0",
    },
    {
        "src": DOCS / "SRS_v3.0_Refactored.md",
        "out": "playIT_SRS_v3.0",
        "title": "Software Requirements Specification (SRS v3.0)",
        "short": "SRS v3.0",
    },
    {
        "src": DOCS / "SPMP_v2.0_Refactored.md",
        "out": "playIT_SPMP_v2.0",
        "title": "Software Project Management Plan (SPMP v2.0)",
        "short": "SPMP v2.0",
    },
]

MVP_FORM = {
    "src": DOCS / "MVP_Validation_Findings_and_Refactoring_Priorities_Filled.md",
    "template": DOCS / "MVP Validation Findings & Refactoring Priorities.docx",
    "filled": DOCS / "MVP_Validation_Findings_and_Refactoring_Priorities_Filled.docx",
    "out": "playIT_MVP_Validation_Findings_and_Refactoring_Priorities",
}

FOOTER_TEXT = "IT411 Capstone & Research 2  ·  CIT-U College of Computer Studies"


# ---------------------------------------------------------------------------
# Spec documents (SDD, SRS, SPMP)
# ---------------------------------------------------------------------------

def _style_font(style, size=None, color=None, bold=None, italic=None):
    style.font.name = BODY_FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in list(rfonts.attrib):
        if attr.endswith("Theme") or attr.endswith("theme"):
            del rfonts.attrib[attr]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)
    if size is not None:
        style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


def setup_styles(doc):
    normal = doc.styles["Normal"]
    _style_font(normal, 10, TEXT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    spacing = {1: (18, 6), 2: (12, 4), 3: (10, 3), 4: (8, 2)}
    colors = {1: TEAL, 2: PLUM, 3: TEAL, 4: TEXT}
    for level in range(1, 5):
        st = doc.styles[f"Heading {level}"]
        _style_font(st, HEADING_SIZES[level], colors[level], bold=True, italic=False)
        st.paragraph_format.space_before = Pt(spacing[level][0])
        st.paragraph_format.space_after = Pt(spacing[level][1])
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.0


def setup_page(section, short_title: str):
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    section.different_first_page_header_footer = True

    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run(f"PlayIT  |  {short_title}")
    set_run_font(run, BODY_FONT, 8, MUTED)
    p_pr = hp._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "4"), ("w:color", "C8D0D8")):
        bottom.set(qn(k), v)
    border.append(bottom)
    p_pr.append(border)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text, field in ((FOOTER_TEXT + "  ·  Page ", "PAGE"), (" of ", "NUMPAGES")):
        set_run_font(fp.add_run(text), BODY_FONT, 8, MUTED)
        add_field(fp, field, "1", font=BODY_FONT, size=8, color=MUTED)


def split_front_matter(lines: list[str]) -> tuple[list[str], list[str]]:
    """The front matter is everything before the first horizontal rule."""
    for idx, line in enumerate(lines):
        if re.match(r"^\s*-{3,}\s*$", line):
            return lines[:idx], lines[idx + 1:]
    return [], lines


def parse_front(front: list[str]) -> dict:
    meta = {"doc_type": "", "project": "", "fields": []}
    for line in front:
        s = line.strip()
        if s.startswith("# "):
            meta["doc_type"] = s[2:].strip()
        elif s.startswith("## "):
            meta["project"] = re.sub(r"^Project:\s*", "", s[3:].strip())
        else:
            m = re.match(r"^\*\*(.+?):\*\*\s*(.*?)\s*$", s)
            if m:
                meta["fields"].append((m.group(1), m.group(2)))
    return meta


def build_cover(doc, meta: dict, title: str):
    fields = dict(meta["fields"])
    place = fields.pop("Department", None) or fields.pop("Institution", None) or ""
    parts = [p.strip() for p in place.split(",")]
    college, institution = (parts[0], ", ".join(parts[1:])) if len(parts) > 1 else ("", place)

    def centered(text, size, color=TEXT, bold=False, before=0, after=4, caps=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        add_inline(p, text.upper() if caps else text, size=size, color=color, bold=bold)
        return p

    centered(institution, 13, TEAL, bold=True, before=36, caps=True)
    if college:
        centered(college, 11, TEXT, after=2)
    if "Degree Program" in fields:
        centered(fields.pop("Degree Program"), 10, MUTED, after=0)

    centered(title, 24, TEAL, bold=True, before=96, after=10)
    project = meta["project"]
    if "—" in project:
        name, subtitle = (s.strip() for s in project.split("—", 1))
        centered(name, 18, PLUM, bold=True, before=6, after=2)
        centered(subtitle, 12.5, PLUM, after=0)
    elif project:
        centered(project, 14, PLUM, bold=True)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(70)
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    for label, value in fields.items():
        row = table.add_row()
        row.cells[0].width, row.cells[1].width = Inches(1.75), Inches(4.5)
        lp, vp = row.cells[0].paragraphs[0], row.cells[1].paragraphs[0]
        for p in (lp, vp):
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
        add_inline(lp, label, size=9.5, color=TEAL, bold=True)
        add_inline(vp, value, size=9.5, color=TEXT)
    page_break(doc)


def build_toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("Table of Contents")
    set_run_font(run, BODY_FONT, 15, TEAL, bold=True)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Open in Word and update fields to build the table of contents.")
    page_break(doc)


def build_spec_document(spec: dict) -> Path:
    lines = spec["src"].read_text(encoding="utf-8").splitlines()
    front, body = split_front_matter(lines)
    doc = docx.Document()
    setup_styles(doc)
    setup_page(doc.sections[0], spec["short"])
    build_cover(doc, parse_front(front), spec["title"])
    build_toc(doc)
    render_blocks(doc, parse_blocks(body), Theme())
    out = SUBMISSION_DIR / f"{spec['out']}.docx"
    doc.save(out)
    print(f"[OK] DOCX  {out.relative_to(REPO)}")
    return out


def add_form_footer(path: Path):
    """Page numbers for the filled course form (the template has an empty footer)."""
    doc = docx.Document(path)
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text, field in (("PlayIT  ·  MVP Validation Findings & Refactoring Priorities  ·  Page ", "PAGE"),
                        (" of ", "NUMPAGES")):
        set_run_font(fp.add_run(text), "Arial", 8, MUTED)
        add_field(fp, field, "1", font="Arial", size=8, color=MUTED)
    doc.save(path)


# ---------------------------------------------------------------------------
# PDF export
# ---------------------------------------------------------------------------

def windows_path(path: Path) -> str:
    path = path.resolve()
    if sys.platform == "win32":
        return str(path)
    try:
        return subprocess.run(["wslpath", "-w", str(path)], capture_output=True, text=True, check=True).stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        m = re.match(r"^/mnt/([a-z])/(.*)$", str(path))
        return f"{m.group(1).upper()}:\\{m.group(2)}".replace("/", "\\") if m else str(path)


def powershell() -> str | None:
    found = shutil.which("powershell.exe") or shutil.which("powershell")
    if found:
        return found
    for candidate in (Path(r"C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe"),
                      Path("/mnt/c/Windows/System32/WindowsPowerShell/v1.0/powershell.exe")):
        if candidate.exists():
            return str(candidate)
    return None


def export_pdf_with_word(docx_path: Path, pdf_path: Path) -> bool:
    """Update the TOC and fields, save the .docx, and export a PDF with heading bookmarks."""
    script = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$w.DisplayAlerts = 0
try {{
    $d = $w.Documents.Open('{windows_path(docx_path)}', $false, $false)
    foreach ($t in $d.TablesOfContents) {{ $null = $t.Update() }}
    $null = $d.Fields.Update()
    foreach ($t in $d.TablesOfContents) {{ $null = $t.UpdatePageNumbers() }}
    $d.Save()
    $d.ExportAsFixedFormat('{windows_path(pdf_path)}', 17, $false, 0, 0, 1, 1, 0, $true, $true, 1, $true, $true, $false)
    $d.Close($false)
    Write-Output 'SUCCESS'
}} finally {{
    $w.Quit()
}}
"""
    shell = powershell()
    if shell is None:
        print("[WARN] PowerShell not found, so Word cannot be used.")
        return False
    try:
        res = subprocess.run([shell, "-NoProfile", "-NonInteractive", "-Command", script],
                             capture_output=True, text=True, timeout=300)
    except (OSError, subprocess.TimeoutExpired) as e:
        print(f"[WARN] Word export failed: {e}")
        return False
    if "SUCCESS" in res.stdout:
        return True
    print(f"[WARN] Word export failed: {res.stderr.strip()[:800]}")
    return False


def export_pdf_with_libreoffice(docx_path: Path, pdf_path: Path) -> bool:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    res = subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent),
                          str(docx_path)], capture_output=True, text=True, timeout=300)
    produced = pdf_path.parent / (docx_path.stem + ".pdf")
    if res.returncode == 0 and produced.exists():
        if produced != pdf_path:
            produced.replace(pdf_path)
        print("[WARN] PDF made with LibreOffice: open the .docx in Word once to fill in the table of contents.")
        return True
    return False


def export_pdf(docx_path: Path) -> Path:
    pdf_path = docx_path.with_suffix(".pdf")
    if export_pdf_with_word(docx_path, pdf_path) or export_pdf_with_libreoffice(docx_path, pdf_path):
        print(f"[OK] PDF   {pdf_path.relative_to(REPO)}")
        return pdf_path
    sys.exit(f"Error: could not convert {docx_path.name} to PDF (needs Microsoft Word or LibreOffice).")


# ---------------------------------------------------------------------------
# Checklist
# ---------------------------------------------------------------------------

CHECKLIST = """# PlayIT — IT411 Submission Package: Overview and Review Checklist
**Course:** IT411 — Capstone & Research 2 | Semester 1, AY 2026–2027
**Degree Program:** Bachelor of Science in Information Technology
**Institution:** College of Computer Studies, Cebu Institute of Technology – University
**Submission Date:** October 10, 2026
**Branch:** `refactor/hear-say-it`. Code validated at `eaf8634` (agy batch validation `c07206e`); package reviewed and rebuilt by Claude on October 10, 2026.

Rebuild everything with `python3 tools/docs/render_submission_package.py` (python-docx and Microsoft Word). The Markdown files in `docs/` are the sources; never edit the .docx or .pdf copies by hand.

---

## 1. Deliverables

| # | Deliverable | Files | Source | Contents |
|:---:|---|---|---|---|
| 1 | **MVP Validation Findings & Refactoring Priorities** (course form) | `playIT_MVP_Validation_Findings_and_Refactoring_Priorities.pdf` / `.docx` | `MVP_Validation_Findings_and_Refactoring_Priorities_Filled.md`, filled into the course template | Sections 1–7 and the declaration: 5 findings, 6 refactoring priorities, summary matrix, non-priority screening |
| 2 | **Software Design Description (SDD v2.0, rev. 2.3)** | `playIT_SDD_v2.0.pdf` / `.docx` | `SDD_v2.0_Refactored.md` | §2.0 implementation status checked against the code on 2026-10-09 |
| 3 | **Software Requirements Specification (SRS v3.0, rev. 3.3)** | `playIT_SRS_v3.0.pdf` / `.docx` | `SRS_v3.0_Refactored.md` | RTM v3.0 (16 findings) and §4.1 automated test coverage |
| 4 | **Software Project Management Plan (SPMP v2.0, rev. 2.3)** | `playIT_SPMP_v2.0.pdf` / `.docx` | `SPMP_v2.0_Refactored.md` | WBS, revised schedule with sprint status, risk register R1–R8 |

## 2. Technical evidence (agy batch validation, `c07206e`)

- **Unit tests:** 364 run, 0 failures, 6 skipped by design (font-scale checks on the 360x640 profile).
- **Roborazzi screenshots:** 41 images: 9 layout checks on 4 sizes (360x640, 360x740, 411x891, 800x1280) plus 5 single-screen captures.
- **Debug APK:** `playit-debug-B-eaf8634.apk` (102 MB). The manifest has no `INTERNET` permission, so the app cannot send data off the device.
- **Zero-emoji policy:** `ZeroEmojiPolicyTest` passes on the app's UI strings. (The course form's own priority options use colored-circle symbols; they are part of the template and left as issued.)

## 3. Before turning in (team)

- [ ] Confirm the adviser's name on the MVP form and the three cover pages (Mr. Joemarie C. Amparo, taken from the team's transmittal notes).
- [ ] Confirm the group name, section ("IT411 G1–G8") and member list on the MVP form.
- [ ] Sign the MVP form's Final Declaration (Team Lead) on the printed copy if the course requires a wet signature.
- [ ] Open each PDF and check that the table of contents, page numbers and landscape matrices (SRS RTM, SPMP risk register) look right.
- [ ] Optional: install the APK on the test phone (`adb install -r playit-debug-B-eaf8634.apk`).

## 4. Still pending after this submission (stated in the documents)

- Gate 3 teacher audit of every shipped clip, and the short-vowel clips (ElevenLabs, due Oct 11).
- Round 2 field validation (Weeks 7–8): ASR agreement, latency, hesitation, completion.
- Planned components: `LessonEngine`, `AudioComposer`, Room schema v4 with telemetry and CSV export, the input-level mic ripple.
"""


def write_checklist():
    path = SUBMISSION_DIR / "SUBMISSION_OVERVIEW_AND_REVIEW_CHECKLIST.md"
    path.write_text(CHECKLIST, encoding="utf-8", newline="\n")
    print(f"[OK] MD    {path.relative_to(REPO)}")


def main():
    SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)
    print("=== Building the CIT-U IT411 submission package ===")

    for spec in SPEC_DOCUMENTS:
        print(f"\n{spec['title']}")
        export_pdf(build_spec_document(spec))

    print("\nMVP Validation Findings & Refactoring Priorities")
    fill_mvp_validation_form.build_form(MVP_FORM["src"], MVP_FORM["template"], MVP_FORM["filled"])
    add_form_footer(MVP_FORM["filled"])
    form_docx = SUBMISSION_DIR / f"{MVP_FORM['out']}.docx"
    shutil.copy2(MVP_FORM["filled"], form_docx)
    print(f"[OK] DOCX  {form_docx.relative_to(REPO)}")
    export_pdf(form_docx)
    shutil.copy2(form_docx, MVP_FORM["filled"])

    print()
    write_checklist()
    print("\n=== Submission package complete ===")


if __name__ == "__main__":
    main()
