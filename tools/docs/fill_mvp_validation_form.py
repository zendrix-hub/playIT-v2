#!/usr/bin/env python3
"""Fill the course's MVP Validation Findings & Refactoring Priorities form.

Source:   docs/MVP_Validation_Findings_and_Refactoring_Priorities_Filled.md
Template: docs/MVP Validation Findings & Refactoring Priorities.docx (as issued by the course)
Output:   docs/MVP_Validation_Findings_and_Refactoring_Priorities_Filled.docx

The template's wording is kept. Each answer replaces the italic placeholder under its
question, every option list becomes a row of ticked or empty boxes, the Finding and
Refactoring Priority blocks are repeated once per item in the Markdown, and both
tables are filled. render_submission_package.py calls build_form() and makes the PDF.
"""

from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from md_to_docx import SYMBOL_FONT, add_inline, set_run_font, split_row

REPO = Path(__file__).resolve().parents[2]
SOURCE = REPO / "docs" / "MVP_Validation_Findings_and_Refactoring_Priorities_Filled.md"
TEMPLATE = REPO / "docs" / "MVP Validation Findings & Refactoring Priorities.docx"
OUTPUT = REPO / "docs" / "MVP_Validation_Findings_and_Refactoring_Priorities_Filled.docx"

ANSWER_SIZE = 10
TABLE_SIZE = 8
GUIDE_COLOR = RGBColor(0x71, 0x80, 0x96)
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"


# ---------------------------------------------------------------------------
# Markdown parsing
# ---------------------------------------------------------------------------

_ITEM = re.compile(r"^(\d+)\.\s+\*\*(.+?)\*\*\s*$")
_LABELED = re.compile(r"^\*\*(.+?):?\*\*:?\s*(.*)$")


def _label_detail(text: str) -> tuple[str, str]:
    """'**Customer:** detail', 'Customer: detail' or 'Customer' -> ('Customer', 'detail')."""
    text = text.strip()
    m = re.match(r"^\*\*(.+?)\*\*\s*(.*)$", text)
    if m:
        label, rest = m.group(1).rstrip(":"), m.group(2).lstrip(":").strip()
        return label.strip(), rest
    if ":" in text:
        label, rest = text.split(":", 1)
        return label.strip(), rest.strip()
    return text, ""


def _sections(lines: list[str]) -> dict[str, list[str]]:
    out, current = {}, None
    for line in lines:
        m = re.match(r"^## (Section \d+|Final Declaration)", line)
        if m:
            current = m.group(1)
            out[current] = []
        elif current:
            if line.strip() == "---":
                continue
            out[current].append(line)
    return out


def _subsections(lines: list[str]) -> list[tuple[str, list[str]]]:
    out = []
    for line in lines:
        if line.startswith("### "):
            out.append((line[4:].strip(), []))
        elif out:
            out[-1][1].append(line)
    return out


def _paragraphs(lines: list[str]) -> list[str]:
    paras, buf = [], []
    for line in lines:
        if line.strip():
            buf.append(line.strip())
        elif buf:
            paras.append(" ".join(buf))
            buf = []
    if buf:
        paras.append(" ".join(buf))
    return paras


def _parse_item_body(lines: list[str]) -> dict:
    item = {"paras": [], "bullets": [], "checks": [], "evidence": []}
    buf = []

    def flush():
        if buf:
            item["paras"].append(" ".join(buf))
            buf.clear()

    for raw in lines:
        line = raw.strip()
        if not line:
            flush()
        elif line.startswith("- [x] ") or line.startswith("- [X] "):
            flush()
            item["checks"].append(_label_detail(line[6:]))
        elif line.startswith("- [ ] "):
            flush()
        elif line.startswith("- "):
            flush()
            item["bullets"].append(line[2:].strip())
        elif line.startswith("*Evidence type:*"):
            flush()
            for part in line[len("*Evidence type:*"):].split(";"):
                if part.strip():
                    item["evidence"].append(_label_detail(part))
        else:
            buf.append(line)
    flush()
    return item


def _parse_blocks(lines: list[str], prefix: str) -> list[dict]:
    blocks = []
    for title, body in _subsections(lines):
        m = re.match(rf"^{prefix} #(\d+):\s*(.+)$", title)
        if not m:
            continue
        items, current = {}, None
        for line in body:
            im = _ITEM.match(line.strip()) if not line.startswith(" ") else None
            if im:
                current = int(im.group(1))
                items[current] = []
            elif current is not None:
                items[current].append(line)
        blocks.append({
            "number": int(m.group(1)),
            "title": m.group(2).strip(),
            "items": {k: _parse_item_body(v) for k, v in items.items()},
        })
    return blocks


def _table(lines: list[str]) -> list[list[str]]:
    rows = [split_row(l) for l in lines if l.strip().startswith("|")]
    return [r for r in rows[2:]]  # drop header and separator


def parse_form_md(path: Path) -> dict:
    lines = path.read_text(encoding="utf-8").splitlines()
    sec = _sections(lines)

    project, members, respondents, current = {}, [], [], None
    for line in sec["Section 1"]:
        s = line.strip()
        if line.startswith("- "):
            label, value = _label_detail(s[2:])
            project[label] = value
            current = label
        elif re.match(r"^\d+\.\s+", s) and current == "Team Members":
            members.append(re.sub(r"^\d+\.\s+", "", s))
        elif s.startswith("- [x] ") and current == "Types of respondents involved":
            respondents.append(_label_detail(s[6:]))
    project["Team Members"] = members

    overview = {}
    for title, body in _subsections(sec["Section 2"]):
        number = int(title.split(".")[0])
        if number == 1:
            overview["purpose"] = _paragraphs(body)
        else:
            fields, last = {}, None
            for line in body:
                s = line.strip()
                if line.startswith("- "):
                    last, value = _label_detail(s[2:])
                    fields[last] = {"value": value, "sub": []}
                elif s and last and line.startswith(" "):
                    fields[last]["sub"].append(re.sub(r"^(\d+\.|-)\s+", "", s))
            overview[number] = fields

    section6 = sec["Section 6"]
    important = []
    for title, body in _subsections(section6):
        if title.startswith("Important Question"):
            important = _paragraphs(body)

    reflection = {}
    for title, body in _subsections(sec["Section 7"]):
        number = int(title.split(".")[0])
        if number == 4:
            reflection[4] = [_label_detail(l.strip()[6:]) for l in body if l.strip().startswith("- [x] ")]
        else:
            reflection[number] = _paragraphs(body)

    declaration = {}
    for line in sec["Final Declaration"]:
        m = re.match(r"^\*\*(.+?):\*\*\s*(.+?)\s*$", line.strip())
        if m:
            declaration[m.group(1)] = m.group(2)

    form = {
        "project": project,
        "respondents": respondents,
        "overview": overview,
        "findings": _parse_blocks(sec["Section 3"], "Finding"),
        "priorities": _parse_blocks(sec["Section 4"], "Refactoring Priority"),
        "matrix": _table(sec["Section 5"]),
        "non_priorities": _table([l for l in section6 if l.strip().startswith("|")]),
        "important": important,
        "reflection": reflection,
        "declaration": declaration,
    }
    _check(form)
    return form


def _check(form: dict):
    for f in form["findings"]:
        missing = set(range(1, 7)) - set(f["items"])
        if missing:
            raise ValueError(f"Finding #{f['number']} is missing questions {sorted(missing)}")
    for p in form["priorities"]:
        missing = set(range(1, 10)) - set(p["items"])
        if missing:
            raise ValueError(f"Refactoring Priority #{p['number']} is missing questions {sorted(missing)}")
    if len(form["non_priorities"]) != 9:
        raise ValueError("Section 6 needs exactly the 9 rows of the form's table")
    for key in ("Team Lead", "Date"):
        if key not in form["declaration"]:
            raise ValueError(f"Final Declaration needs **{key}:**")


# ---------------------------------------------------------------------------
# Template editing helpers
# ---------------------------------------------------------------------------

def _clear_runs(p: Paragraph):
    for child in list(p._p):
        if child.tag in (qn("w:r"), qn("w:hyperlink")):
            p._p.remove(child)
    p_pr = p._p.pPr
    if p_pr is not None:
        r_pr = p_pr.find(qn("w:rPr"))
        if r_pr is not None:
            p_pr.remove(r_pr)


def _truncate_at_break(p: Paragraph):
    """Keep the bold question and drop the line break and hint that follow it."""
    found = False
    for run in list(p._p.findall(qn("w:r"))):
        if found:
            p._p.remove(run)
            continue
        br = run.find(qn("w:br"))
        if br is not None:
            for el in list(run):
                if found:
                    run.remove(el)
                elif el is br:
                    run.remove(el)
                    found = True


def _new_paragraph_after(anchor, body) -> Paragraph:
    el = OxmlElement("w:p")
    anchor.addnext(el)
    p = Paragraph(el, body)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    return p


def _write_answer(p: Paragraph, text: str):
    add_inline(p, text, size=ANSWER_SIZE)


def _guide(p: Paragraph):
    """Show template instructions as small grey guidance text."""
    for run in p.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = GUIDE_COLOR


def _bullet(p: Paragraph, text: str, left_twips=720):
    p_pr = p._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left_twips))
    ind.set(qn("w:hanging"), "280")
    p_pr.append(ind)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("•\t"), size=ANSWER_SIZE)
    add_inline(p, text, size=ANSWER_SIZE)


def _option_key(text: str) -> str:
    text = re.sub(r"^[^\w]+", "", text.strip())
    text = re.split(r"\s+—\s+|:", text)[0]
    return text.strip().lower()


def _set_checkbox(p: Paragraph, checked: bool, detail: str = "", detail_on_new_line=False):
    """Turn a template option (a bulleted list item) into a ticked or empty box."""
    label = p.text.strip()
    p_pr = p._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    level = 0
    if num_pr is not None:
        ilvl = num_pr.find(qn("w:ilvl"))
        level = int(ilvl.get(qn("w:val"))) if ilvl is not None else 0
        p_pr.remove(num_pr)
    for old in p_pr.findall(qn("w:ind")):
        p_pr.remove(old)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(720 * (level + 1)))
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    _clear_runs(p)
    glyph = p.add_run("☒\t" if checked else "☐\t")
    set_run_font(glyph, SYMBOL_FONT, ANSWER_SIZE + 1)
    if _option_key(label) == "other" and checked:
        label = "Other:"
    if checked:
        run = p.add_run(label)
        set_run_font(run, size=ANSWER_SIZE, bold=True)
        if detail:
            if detail_on_new_line:
                p.add_run().add_break()
            else:
                p.add_run(" " if label.endswith(":") else ": ")
            add_inline(p, detail, size=ANSWER_SIZE)
    else:
        set_run_font(p.add_run(label), size=ANSWER_SIZE)


def _fill_options(options: list[Paragraph], checks: list[tuple[str, str]], what: str, detail_on_new_line=False):
    keys = {_option_key(p.text): p for p in options}
    chosen = {}
    for label, detail in checks:
        key = _option_key(label)
        if key not in keys:
            raise ValueError(f"'{label}' is not an option of '{what}'. Options: {[p.text for p in options]}")
        chosen[key] = detail
    for p in options:
        key = _option_key(p.text)
        _set_checkbox(p, key in chosen, chosen.get(key, ""), detail_on_new_line)


def _replace_with_answer(p: Paragraph, paras: list[str], body):
    """Replace an italic placeholder paragraph with the answer paragraph(s)."""
    _clear_runs(p)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    if not paras:
        return p
    _write_answer(p, paras[0])
    last = p
    for text in paras[1:]:
        last = _new_paragraph_after(last._p, body)
        _write_answer(last, text)
    return last


def _insert_answers_after(anchor: Paragraph, item: dict, body, left_twips=720) -> Paragraph:
    last = anchor
    for text in item["paras"]:
        last = _new_paragraph_after(last._p, body)
        _write_answer(last, text)
    for text in item["bullets"]:
        last = _new_paragraph_after(last._p, body)
        _bullet(last, text, left_twips)
    return last


def _strip_ids(el):
    for node in el.iter():
        for attr in (W14 + "paraId", W14 + "textId"):
            if attr in node.attrib:
                del node.attrib[attr]


def _set_heading(p: Paragraph, text: str):
    _clear_runs(p)
    run = p.add_run(text)
    run.bold = True


def _body_paragraphs(doc) -> list[Paragraph]:
    return [Paragraph(el, doc._body) for el in doc.element.body if el.tag == qn("w:p")]


def _find(paras: list[Paragraph], prefix: str, start: int = 0) -> int:
    for idx in range(start, len(paras)):
        if paras[idx].text.strip().startswith(prefix):
            return idx
    raise ValueError(f"Template paragraph starting with '{prefix}' not found")


def _following_options(paras: list[Paragraph], idx: int) -> list[Paragraph]:
    """The bulleted option list right after paragraph idx."""
    out = []
    for p in paras[idx + 1:]:
        if p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None:
            out.append(p)
        elif out:
            break
        elif p.text.strip() in ("", "Examples:"):
            continue
        else:
            break
    return out


# ---------------------------------------------------------------------------
# Section fillers
# ---------------------------------------------------------------------------

def _fill_section1(doc, form):
    paras = _body_paragraphs(doc)
    project = form["project"]
    for label in ("Team/Group Name", "Project/System Title", "Program / Section", "Team Members", "Adviser",
                  "System URL / MVP Link", "Date of MVP Validation", "Number of respondents/participants"):
        p = paras[_find(paras, label)]
        value = project[label]
        if isinstance(value, list):
            p.add_run(":")
            for member in value:
                p.add_run().add_break()
                add_inline(p, member, size=ANSWER_SIZE)
        else:
            p.add_run(": ")
            add_inline(p, value, size=ANSWER_SIZE)
    idx = _find(paras, "Types of respondents involved")
    paras[idx].add_run(":")
    _fill_options(_following_options(paras, idx), form["respondents"], "Types of respondents involved")


def _fill_section2(doc, form):
    body = doc._body
    paras = _body_paragraphs(doc)
    ov = form["overview"]

    q1 = _find(paras, "1. What was the primary purpose")
    placeholder, instruction = paras[q1 + 1], paras[q1 + 2]
    paras[q1]._p.addnext(instruction._p)
    _guide(instruction)
    _replace_with_answer(placeholder, ov["purpose"], body)

    for label, key in (("Framework/Model", "Framework/Model"),
                       ("Key constructs/criteria evaluated", "Key constructs/criteria evaluated"),
                       ("Why was this framework/model appropriate", "Why was this framework/model appropriate for your project?")):
        p = paras[_find(paras, label)]
        field = ov[2][key]
        template_label = p.text.split("__")[0].strip()
        _clear_runs(p)
        set_run_font(p.add_run(template_label), bold=True)
        p.add_run(" " if template_label.endswith(("?", ":")) else ": ")
        if field["value"]:
            add_inline(p, field["value"], size=ANSWER_SIZE)
        for n, sub in enumerate(field["sub"], start=1):
            p.add_run().add_break()
            add_inline(p, f"{n}. {sub}", size=ANSWER_SIZE)

    include = paras[_find(paras, "Include:")]
    _guide(include)
    for label, key in (("Who participated", "Who participated"),
                       ("How they interacted with the MVP", "How they interacted with the MVP"),
                       ("What activities/tasks they performed", "What activities/tasks they performed"),
                       ("How feedback/data was collected", "How feedback/data was collected")):
        p = paras[_find(paras, label)]
        _clear_runs(p)
        set_run_font(p.add_run(label + ": "), bold=True)
        add_inline(p, ov[3][key]["value"], size=ANSWER_SIZE)


def _fill_finding(block: list[Paragraph], finding: dict, body):
    items = finding["items"]
    _set_heading(block[0], f"Finding #{finding['number']}: {finding['title']}")
    q = {n: next(i for i, p in enumerate(block) if p.text.strip().startswith(f"{n}. ")) for n in range(1, 7)}

    _truncate_at_break(block[q[1]])
    _insert_answers_after(block[q[1]], items[1], body)

    _truncate_at_break(block[q[2]])
    last = _insert_answers_after(block[q[2]], items[2], body)
    label = _new_paragraph_after(last._p, body)
    label.paragraph_format.space_after = Pt(2)
    label.add_run("Type of evidence (examples from the form):")
    _guide(label)
    _fill_options(block[q[2] + 1:q[3]], items[2]["evidence"], "evidence")

    _fill_options(block[q[3] + 1:q[4]], items[3]["checks"], "stakeholder")
    _replace_with_answer(block[q[4] + 1], items[4]["paras"], body)
    _fill_options(block[q[5] + 1:q[6]], items[5]["checks"], "measurable outcome")
    _replace_with_answer(block[q[6] + 1], items[6]["paras"], body)


def _fill_priority(block: list[Paragraph], prio: dict, body):
    items = prio["items"]
    _set_heading(block[0], f"Refactoring Priority #{prio['number']}: {prio['title']}")
    q = {n: next(i for i, p in enumerate(block) if p.text.strip().startswith(f"{n}. ")) for n in range(1, 10)}
    for n in (1, 2, 4, 5, 6, 7):
        _replace_with_answer(block[q[n] + 1], items[n]["paras"], body)
    _fill_options(block[q[3] + 1:q[4]], items[3]["checks"], "type of refactoring")

    examples = block[q[8] + 1]
    _insert_answers_after(block[q[8]], items[8], body)
    _guide(examples)
    _fill_options(block[q[8] + 2:q[9]], items[8]["checks"], "measurement")
    _fill_options([p for p in block[q[9] + 1:] if p.text.strip()], items[9]["checks"], "priority",
                  detail_on_new_line=True)


def _clone_blocks(doc, start_prefix: str, end_prefix: str, count: int):
    """Copy the template block [start, end) count times in place; returns each copy's paragraphs."""
    paras = _body_paragraphs(doc)
    start, end = _find(paras, start_prefix), _find(paras, end_prefix)
    proto = [p._p for p in paras[start:end]]
    anchor = proto[0].getprevious()
    copies = []
    for _ in range(count):
        els = [copy.deepcopy(el) for el in proto]
        for el in els:
            _strip_ids(el)
            anchor.addnext(el)
            anchor = el
        copies.append([Paragraph(el, doc._body) for el in els])
    for el in proto:
        el.getparent().remove(el)
    return copies


def _fill_section3(doc, form):
    paras = _body_paragraphs(doc)
    # The template has "Finding #2 / #3: Repeat the same fields." after the Finding #1 block.
    first_repeat = _find(paras, "Finding #2")
    recommended = _find(paras, "Recommended:")
    for p in paras[first_repeat:recommended]:
        if p.text.strip():
            p._p.getparent().remove(p._p)
    for block, finding in zip(_clone_blocks(doc, "Finding #1", "Recommended:", len(form["findings"])),
                              form["findings"]):
        _fill_finding(block, finding, doc._body)


def _fill_section4(doc, form):
    copies = _clone_blocks(doc, "Refactoring Priority #1", "Section 5", len(form["priorities"]))
    for block, prio in zip(copies, form["priorities"]):
        _fill_priority(block, prio, doc._body)


def _fill_cell(cell, text: str, bold=False):
    p = cell.paragraphs[0]
    _clear_runs(p)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    add_inline(p, text, size=TABLE_SIZE, bold=bold, break_code=True)


def _set_widths(table, inches: list[float], header_size: float):
    """Widen the template's columns so no header word breaks, and keep the header on each page."""
    grid = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for col, width in zip(grid, inches):
        col.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        for cell, width in zip(row.cells, inches):
            cell.width = Inches(width)
    for p in table.rows[0].cells[0]._tc.getparent().iter(qn("w:p")):
        for r in Paragraph(p, table).runs:
            r.font.size = Pt(header_size)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _fill_section5(doc, form):
    table = doc.tables[0]
    rows = form["matrix"]
    while len(table.rows) - 1 < len(rows):
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
    for r_idx, values in enumerate(rows, start=1):
        for c_idx, value in enumerate(values):
            _fill_cell(table.rows[r_idx].cells[c_idx], value, bold=(c_idx == 0))
    _set_widths(table, [0.35, 1.05, 1.0, 1.05, 1.05, 1.05, 0.75], header_size=9)


def _fill_section6(doc, form):
    table = doc.tables[1]
    for r_idx, (item, identified, justification) in enumerate(form["non_priorities"], start=1):
        row = table.rows[r_idx]
        if _option_key(row.cells[0].text) != _option_key(item):
            raise ValueError(f"Section 6 row {r_idx}: '{item}' does not match the form's '{row.cells[0].text}'")
        for t in row.cells[0]._tc.iter(qn("w:t")):
            t.text = t.text.replace("/", "/\u200b")  # let the long "a/b/c/d" label wrap at slashes
        _fill_cell(row.cells[1], identified, bold=True)
        action = row.cells[2]
        p = action.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        add_inline(p, justification, size=TABLE_SIZE)
    _set_widths(table, [2.35, 0.95, 3.2], header_size=10)

    paras = _body_paragraphs(doc)
    instruction = paras[_find(paras, "Provide evidence-based justification")]
    _guide(instruction)
    last = instruction
    for text in form["important"]:
        last = _new_paragraph_after(last._p, doc._body)
        _write_answer(last, text)


def _fill_section7(doc, form):
    paras = _body_paragraphs(doc)
    refl = form["reflection"]
    start = _find(paras, "Section 7")
    for n in (1, 2, 3):
        q = _find(paras, f"{n}. ", start)
        _replace_with_answer(paras[q + 1], refl[n], doc._body)
    q4 = _find(paras, "4. What requirements and/or design documents", start)
    _fill_options(_following_options(paras, q4), refl[4], "documents to update")


def _fill_declaration(doc, form):
    paras = _body_paragraphs(doc)
    p = paras[_find(paras, "Team Lead:")]
    _clear_runs(p)
    set_run_font(p.add_run("Team Lead: "), bold=True)
    add_inline(p, form["declaration"]["Team Lead"], size=ANSWER_SIZE + 1)
    p.add_run().add_break()
    set_run_font(p.add_run("Date: "), bold=True)
    add_inline(p, form["declaration"]["Date"], size=ANSWER_SIZE + 1)


def build_form(source: Path = SOURCE, template: Path = TEMPLATE, output: Path = OUTPUT) -> Path:
    form = parse_form_md(source)
    doc = docx.Document(template)
    _fill_section1(doc, form)
    _fill_section2(doc, form)
    _fill_section3(doc, form)
    _fill_section4(doc, form)
    _fill_section5(doc, form)
    _fill_section6(doc, form)
    _fill_section7(doc, form)
    _fill_declaration(doc, form)
    doc.save(output)
    print(f"[OK] DOCX  {output.relative_to(REPO)} "
          f"({len(form['findings'])} findings, {len(form['priorities'])} refactoring priorities)")
    return output


if __name__ == "__main__":
    sys.exit(0 if build_form() else 1)
